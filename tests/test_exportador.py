from controllers.exportador import Exportador


def test_gerar_txt_ofertas_vazio():
    success, validas, descartadas = Exportador.gerar_txt_ofertas([], "dummy_path.txt")
    assert not success
    assert validas == 0
    assert descartadas == 0

def test_gerar_txt_ofertas_descarte_negativos_e_palavras_bloqueadas(tmp_path):
    transacoes = [
        {"valor": -50.00, "descricao": "TARIFA BANCARIA", "data": "12/05/2026"},
        {"valor": 150.00, "descricao": "APLICACAO AUTOMATICA", "data": "12/05/2026"},
        {"valor": 300.00, "descricao": "TRANSFERENCIA RECEBIDA", "data": "13/05/2026"},
        {"valor": 200.00, "descricao": "OFERTA DE CULTO ESP", "data": "14/05/2026"}
    ]

    filepath = tmp_path / "ofertas_teste.txt"
    success, validas, descartadas = Exportador.gerar_txt_ofertas(transacoes, str(filepath))

    # -50.00 (negativo), APLICACAO (palavra bloqueada), TRANSFERENCIA (palavra bloqueada) -> descartados (3)
    # OFERTA DE CULTO ESP -> válida (1)
    assert success
    assert validas == 1
    assert descartadas == 3

    # Verifica o conteúdo do arquivo
    content = filepath.read_text(encoding='utf-8')
    assert "14/05/2026;OFERTA DE CULTO ESP;200,00" in content

def test_gerar_txt_ofertas_usa_historico_e_filtra_pelo_memo(tmp_path):
    transacoes = [
        # Nome do pagador contém "APLICA", mas o filtro olha só o MEMO: deve entrar
        {"valor": 100.00, "descricao": "PIX RECEBIDO - OUTRA IF",
         "historico": "PIX RECEBIDO - OUTRA IF - Maria Aplicada", "data": "03/08/2026"},
        # Depósito em dinheiro continua bloqueado
        {"valor": 80.00, "descricao": "DEPOSITO EM DINHEIRO AG",
         "historico": "DEPOSITO EM DINHEIRO AG", "data": "05/08/2026"},
    ]

    filepath = tmp_path / "ofertas_sicoob.txt"
    success, validas, descartadas = Exportador.gerar_txt_ofertas(transacoes, str(filepath))

    assert success
    assert validas == 1
    assert descartadas == 1
    content = filepath.read_text(encoding="utf-8")
    assert "03/08/2026;PIX RECEBIDO - OUTRA IF - Maria Aplicada;100,00" in content

def test_gerar_excel_lote(tmp_path):
    telemetria = {
        "total_contas": 2,
        "ofx_itens": 10,
        "siga_itens": 8,
        "injecoes": 2,
        "pendentes": 1,
        "volume_financeiro": 1500.50
    }

    dados_lote = [
        {
            "tipo_adm": "ADMINISTRACAO",
            "nome_adm": "CONTA PRINCIPAL",
            "tipo_extrato": "CONTA CORRENTE",
            "conta_id": "12345-6",
            "transacoes": [{"valor": 100.00}],
            "extrato_siga": [{"valor": 100.00}],
            "pendentes": []
        }
    ]

    filepath = tmp_path / "lote_resumo.xlsx"
    Exportador.gerar_excel_lote(telemetria, dados_lote, str(filepath))

    assert filepath.exists()
    assert filepath.stat().st_size > 0


def test_filtrar_ofertas_mantem_so_positivos_sem_palavra_bloqueada():
    transacoes = [
        {"valor": -10.0, "descricao": "TARIFA", "data": "01/08/2026"},
        {"valor": 50.0, "descricao": "RESGATE POUPANCA", "data": "02/08/2026"},
        {"valor": 30.0, "descricao": "PIX RECEBIDO", "data": "03/08/2026"},
    ]
    ofertas, descartadas = Exportador.filtrar_ofertas(transacoes)
    assert [tx["valor"] for tx in ofertas] == [30.0]
    assert descartadas == 2


def test_resumo_ofertas_usa_ultimo_dia_do_mes_mais_recente():
    ofertas = [
        {"valor": 10.10, "data": "15/01/2026"},
        {"valor": 20.20, "data": "03/02/2026"},
        {"valor": 0.05, "data": "data invalida"},
    ]
    total, data_importacao, competencia = Exportador.resumo_ofertas(ofertas)
    assert total == 30.35
    assert data_importacao == "28/02/2026"
    assert competencia == "02/2026"


def test_resumo_ofertas_sem_datas_validas():
    assert Exportador.resumo_ofertas([{"valor": 5.0, "data": ""}]) == (5.0, "", "")


def test_formatar_moeda_no_padrao_do_siga():
    assert Exportador.formatar_moeda(1120441.6) == "R$ 1.120.441,60"
    assert Exportador.formatar_moeda(5.5) == "R$ 5,50"


def test_excel_do_lote_tem_a_aba_prova_real_com_as_sobras(tmp_path):
    import openpyxl

    from controllers.prova_real import MovimentoSiga, ResultadoExtrato, conferir_saldo

    resultados = [
        ResultadoExtrato("12345-6", "10130", "DR - CENTRAL", "CORRENTE", "conciliado", lancados=1,
                         saldo=conferir_saldo(1.0, 1.0), sobras=[]),
        ResultadoExtrato("12345-6", "10396", "DR - CENTRAL", "APLICACAO", "conciliado",
                         saldo=conferir_saldo(50.0, 92.1), sobras=[MovimentoSiga("05/05/2026", 42.1, "DUPLICADO")]),
    ]
    filepath = tmp_path / "lote.xlsx"

    Exportador.gerar_excel_lote({"extratos": resultados}, [], str(filepath))

    ws = openpyxl.load_workbook(filepath)["Prova Real"]
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0][:6] == ("Localidade", "Conta SIGA", "Conta Bancária", "Saldo Banco", "Saldo SIGA", "Diferença")
    assert linhas[1][5] is None and linhas[1][8].startswith("✔ Saldo confere")
    assert linhas[2][5] == 42.1
    assert ("DR - CENTRAL", "10396", "05/05/2026", "DUPLICADO", 42.1) in [tuple(linha[:5]) for linha in linhas]


def test_excel_do_lote_tem_a_aba_baixas_de_energia(tmp_path):
    import openpyxl

    from controllers.baixa_energia import (
        BAIXADA,
        BAIXAR,
        SEM_FATURA,
        DebitoEnergia,
        Decisao,
        ProvisaoSiga,
        ResultadoBaixa,
    )
    from controllers.prova_real import ResultadoExtrato
    from models.indice_faturas_energia import FaturaEnergia

    fatura = FaturaEnergia("018000001", "1.234.567.051-00", 224.25, "26/08/2026", "00012345678")
    baixada = Decisao(DebitoEnergia("26/08/2026", 224.25, "00012345678"), BAIXAR, "", fatura,
                      ProvisaoSiga("018000001", 224.25, "Lançado", "G"))
    sem_fatura = Decisao(DebitoEnergia("26/08/2026", 10.0, "00099999999"), SEM_FATURA, "código sem fatura")
    resultados = [ResultadoExtrato("12345-6", "10123", "ADM - CENTRAL", "CORRENTE", "conciliado",
                                   baixas=[ResultadoBaixa(baixada, BAIXADA), ResultadoBaixa(sem_fatura, SEM_FATURA)])]
    filepath = tmp_path / "lote.xlsx"

    Exportador.gerar_excel_lote({"extratos": resultados}, [], str(filepath))

    linhas = list(openpyxl.load_workbook(filepath)["Baixas de Energia"].iter_rows(values_only=True))
    assert linhas[1][:7] == ("ADM - CENTRAL", "10123", "26/08/2026", 224.25, "00012345678", "1.234.567.051-00",
                             "018000001")
    assert "aprovada e baixada" in linhas[1][7]
    assert linhas[2][6] is None and "código sem fatura" in linhas[2][7]


def test_excel_sem_baixas_nao_cria_a_aba(tmp_path):
    import openpyxl

    from controllers.prova_real import ResultadoExtrato

    filepath = tmp_path / "lote.xlsx"
    Exportador.gerar_excel_lote({"extratos": [ResultadoExtrato("1", "10130", "DR", "CORRENTE", "conciliado")]}, [],
                                str(filepath))
    assert "Baixas de Energia" not in openpyxl.load_workbook(filepath).sheetnames
